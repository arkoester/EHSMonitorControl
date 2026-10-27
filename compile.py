#!/usr/bin/env python3
"""
EHS Room Displays - compiler (GitHub Actions edition).

Reads the registry workbook and the scope & sequence files in src/, scrapes the
school's Daily Announcements page, and writes data.json + version.txt beside the
site. Screens pick the files up on their own - no site edits, ever.

Runs on a schedule inside GitHub Actions (.github/workflows/publish.yml), so no
PC has to be on. "Publish Now" = Actions tab > Publish displays > Run workflow.

Layout (all inside the repo):
  src/registry/EHS_room_display_data.xlsx   Settings, Weeks, Bell, DayTypes, Rooms
  src/sns/*.xlsx                             scope & sequence files; filenames do not
                                             matter - the course name is read from
                                             row 1 of each file
  src/day-types.txt                          date -> day type overrides (half days,
                                             surprise closures), editable from a phone
  data.json, version.txt                     what the screens read (written here)
  compile-report.txt                         what happened last run (written here)

Safety net: a course that fails to compile keeps yesterday's rows; a failed
announcements scrape keeps yesterday's announcements; a fatal registry error
publishes nothing and exits 1 (previous files untouched).

Usage:
  python3 compile.py [--site DIR] [--src DIR] [--no-fetch] [--html FILE]
    --no-fetch   skip the announcements scrape (keeps previous announcements)
    --html FILE  parse announcements from a saved HTML file instead of the web
"""
import argparse
import datetime as dt
import html as htmlmod
import json
import os
import re
import sys
import urllib.request

try:
    import openpyxl
except ImportError:
    print("STOP: openpyxl is not installed (pip install openpyxl)")
    sys.exit(1)

ANNOUNCEMENTS_URL = "https://ehs.ecusd7.org/link-2"
REGISTRY_FILE = "EHS_room_display_data.xlsx"
USER_AGENT = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) EHS-Displays-Compiler"
CENTRAL = None
try:
    from zoneinfo import ZoneInfo
    CENTRAL = ZoneInfo("America/Chicago")
except Exception:  # pragma: no cover
    CENTRAL = None


def now_local():
    return dt.datetime.now(CENTRAL) if CENTRAL else dt.datetime.now()


issues, notes = [], []


# ---------------------------------------------------------------- helpers
def norm_key(h):
    return re.sub(r"[^A-Za-z0-9]+", "_", str(h or "")).strip("_").lower()


def fmt_cell(v):
    if v is None:
        return ""
    if isinstance(v, dt.datetime):
        if v.hour == 0 and v.minute == 0 and v.second == 0:
            return v.strftime("%Y-%m-%d")
        return v.strftime("%-I:%M %p")
    if isinstance(v, dt.date):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, dt.time):
        return v.strftime("%-I:%M %p")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def read_tab(wb, sheet):
    """Registry tabs: a note in row 1, headers in row 2, data from row 3."""
    if sheet not in wb.sheetnames:
        raise KeyError("tab '%s' not found (tabs: %s)" % (sheet, ", ".join(wb.sheetnames)))
    ws = wb[sheet]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    if not rows:
        return []
    headers = [norm_key(h) for h in rows[0]]
    out = []
    for r in rows[1:]:
        o, any_ = {}, False
        for k, v in zip(headers, r):
            if not k:
                continue
            s = fmt_cell(v)
            o[k] = s
            if s:
                any_ = True
        if any_:
            out.append(o)
    return out


def parse_date(s):
    s = str(s or "").strip()
    for f in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y", "%Y-%m-%d %H:%M:%S"):
        try:
            return dt.datetime.strptime(s, f).date()
        except ValueError:
            pass
    m = re.match(r"(\d{1,2})[/-](\d{1,2})[/-](\d{2,4})", s)
    if m:
        y = int(m.group(3))
        y = y + 2000 if y < 100 else y
        return dt.date(y, int(m.group(1)), int(m.group(2)))
    raise ValueError("unreadable date '%s'" % s)


MONTHS = {"jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6, "jul": 7, "aug": 8, "sep": 9, "oct": 10, "nov": 11, "dec": 12}
MONTH_RE = re.compile(r"\b(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\b", re.I)


def month_start(tf, mo, first_day):
    """First Monday of that month in the school year; 'mid' -> second, 'late' -> third."""
    yr = first_day.year if mo >= 7 else first_day.year + 1
    d = dt.date(yr, mo, 1)
    while d.weekday() != 0:
        d += dt.timedelta(days=1)
    if re.search(r"\bmid", tf, re.I):
        d += dt.timedelta(days=7)
    elif re.search(r"\blate", tf, re.I):
        d += dt.timedelta(days=14)
    if d < first_day:
        d = first_day
    return d.isoformat()


def test_current(v):
    return bool(re.match(r"^(x|yes|y|true|1|now|current)$", str(v or "").strip(), re.I))


def normalize_course(t):
    t = re.sub(r"\([^)]*\)", " ", str(t or ""))
    t = re.split(r"\s+[-–—]\s+", t)[0]
    return re.sub(r"[^A-Za-z0-9]", "", t).lower()


def course_words(t):
    t = re.sub(r"\([^)]*\)", " ", str(t or ""))
    t = re.split(r"\s+[-–—]\s+", t)[0]
    return [w for w in re.split(r"[^a-z0-9]+", t.lower()) if len(w) >= 2 and w not in ("and", "the", "of", "to", "in", "for")]


def honorific(teacher):
    t = str(teacher or "").strip()
    m = re.match(r"^\s*([^,]+),\s*(\S)", t)
    return "%s. %s" % (m.group(2), m.group(1)) if m else t


# ---------------------------------------------------------------- scope & sequence
def cell_text(ws, addr):
    v = ws[addr].value
    return "" if v is None else str(v)


def read_sns_course(path, week_starts, first_day, anchors):
    """District S&S template: unit blocks from row 3 down. A unit starts on every row
    whose Chapter/Unit cell (B) has text - the template merges B across its 12-row
    block, but some departments use 2-row blocks, so the stride is not assumed.
    A = Time Frame, B = Chapter/Unit (first line = wall title), C = Learning Target(s),
    H = Notes ('x' / 'now' / 'current' at the start = pacing override)."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    ws = wb.worksheets[0]
    units = []
    for r in range(3, min(ws.max_row, 400) + 1):
        unit_cell = cell_text(ws, "B%d" % r)
        if not unit_cell.strip():
            continue
        if not cell_text(ws, "C%d" % r).strip() and not cell_text(ws, "A%d" % r).strip():
            continue  # stray text in B with no time frame and no targets is not a unit
        unit = unit_cell.split("\n")[0].strip()
        tf = cell_text(ws, "A%d" % r).replace("\n", " ")
        note_cell = cell_text(ws, "H%d" % r)
        lt, sc = [], []
        for line in cell_text(ws, "C%d" % r).split("\n"):
            line = line.strip()
            if not line:
                continue
            (sc if re.match(r"^(PF:\s*)?I can", line) else lt).append(line)
        if not lt:
            lt, sc = sc, []
        starts = ""
        m = re.search(r"[Ww]eeks?\s*(\d+)", tf)
        if m:
            wk = int(m.group(1))
            if wk in week_starts:
                starts = week_starts[wk]
            elif first_day:
                starts = (first_day + dt.timedelta(days=7 * (wk - 1))).isoformat()
            mm = MONTH_RE.search(tf)
            if starts and first_day and mm:
                mo = MONTHS[mm.group(1).lower()]
                if parse_date(starts).month != mo:
                    starts = month_start(tf, mo, first_day)
        elif re.search(r"(\d{1,2})[/-](\d{1,2})[/-](\d{2,4})", tf):
            starts = parse_date(re.search(r"(\d{1,2}[/-]\d{1,2}[/-]\d{2,4})", tf).group(1)).isoformat()
        elif re.search(r"\b(q|quarter\s*)([1-4])\b", tf, re.I) and ("q" + re.search(r"\b(q|quarter\s*)([1-4])\b", tf, re.I).group(2)) in anchors:
            starts = anchors["q" + re.search(r"\b(q|quarter\s*)([1-4])\b", tf, re.I).group(2)]
        elif re.search(r"\b(s|sem(ester)?\s*)([12])\b", tf, re.I) and ("s" + re.search(r"\b(s|sem(ester)?\s*)([12])\b", tf, re.I).group(3)) in anchors:
            starts = anchors["s" + re.search(r"\b(s|sem(ester)?\s*)([12])\b", tf, re.I).group(3)]
        elif first_day and MONTH_RE.search(tf):
            starts = month_start(tf, MONTHS[MONTH_RE.search(tf).group(1).lower()], first_day)
        cur = "x" if re.match(r"^\s*(x|now|current)\b", note_cell, re.I) else ""
        units.append({"unit": unit, "starts": starts, "learning_target_s": "\n".join(lt), "success_criteria": "\n".join(sc), "current": cur, "notes": ""})
    # keep units in sequence: a unit that resolves on/before the previous one starts a week later
    prev = None
    for u in units:
        if not u["starts"]:
            continue
        d = parse_date(u["starts"])
        if prev and d <= prev:
            d = prev + dt.timedelta(days=7)
            u["starts"] = d.isoformat()
        prev = d
    return units


def build_sns_index(folder):
    """Every .xlsx under src/sns is opened and asked which course it is (row 1)."""
    index, out_notes = {}, []
    if not os.path.isdir(folder):
        return index, out_notes
    for root, _dirs, files in os.walk(folder):
        for fn in sorted(files):
            if not fn.lower().endswith(".xlsx") or fn.startswith("~$") or fn.startswith("_"):
                continue
            full = os.path.join(root, fn)
            name = ""
            try:
                wb = openpyxl.load_workbook(full, data_only=True, read_only=False)
                ws = wb.worksheets[0]
                for addr in ("C1", "D1", "B1", "E1", "C2", "B2"):
                    t = cell_text(ws, addr).strip()
                    if t and not re.match(r"^(course\s*name|last\s*update|time\s*frame)", t, re.I):
                        name = t
                        break
            except Exception as e:  # noqa: BLE001
                out_notes.append("sns/%s: could not be opened (%s)." % (fn, e))
                continue
            if not name:
                out_notes.append("sns/%s: no course name found in row 1 - put the course name next to 'Course Name:'." % fn)
                continue
            k = normalize_course(name)
            if k in index:
                out_notes.append("sns/%s: duplicate course '%s' (already provided by %s) - using the first." % (fn, name, os.path.basename(index[k]["file"])))
                continue
            index[k] = {"file": full, "label": name}
    return index, out_notes


# ---------------------------------------------------------------- announcements
def parse_announcements(html):
    """Daily Announcements page: a bold 'ANNOUNCEMENTS <date>' header, then one <p>
    per item, each led by a bold phrase. Returns [{date,title,text}]."""
    start = html.find("fsPageContent")
    body = html[start if start >= 0 else 0:]
    stop = body.find("fsFooter")
    if stop > 0:
        body = body[:stop]
    date, items = "", []
    for m in re.finditer(r"<p[^>]*>(.*?)</p>", body, re.S):
        inner = m.group(1)
        lm = re.search(r"<(?:strong|b)[^>]*>(.*?)</(?:strong|b)>", inner, re.S)
        lead = ""
        if lm:
            lead = htmlmod.unescape(re.sub(r"<[^>]+>", " ", lm.group(1)))
            lead = re.sub(r"\s+", " ", lead).strip()
        flat = htmlmod.unescape(re.sub(r"<[^>]+>", " ", inner))
        flat = re.sub(r"\s+", " ", flat).strip()
        if not flat:
            continue
        dm = re.match(r"^ANNOUNCEMENTS\s+(.+)$", lead)
        if dm:
            date = dm.group(1).strip()
            continue
        if len(lead) > 140 or len(flat) > 1500:
            continue
        if not lead:
            if not date or len(flat) < 20:
                continue
            cut = re.match(r"^(.{10,70}?[.!?])(\s+|$)", flat)
            lead = cut.group(1).strip() if cut else flat[:70].strip()
        rest = flat
        if flat.startswith(lead):
            rest = flat[len(lead):].lstrip(" ,:;-.")
        items.append({"date": date, "title": lead, "text": rest})
    return items


def fetch_announcements(url):
    req = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
    with urllib.request.urlopen(req, timeout=30) as resp:
        return resp.read().decode("utf-8", "replace")


# ---------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--site", default=".")
    ap.add_argument("--src", default=None)
    ap.add_argument("--no-fetch", action="store_true")
    ap.add_argument("--html", default=None)
    args = ap.parse_args()
    site = os.path.abspath(args.site)
    src = os.path.abspath(args.src or os.path.join(site, "src"))
    data_path = os.path.join(site, "data.json")
    report_path = os.path.join(site, "compile-report.txt")
    started = now_local()
    print("Run started %s  cwd=%s" % (started.strftime("%Y-%m-%d %H:%M %Z"), os.getcwd()))

    # previous compile (per-course safety net)
    prev = None
    if os.path.exists(data_path):
        try:
            with open(data_path, encoding="utf-8-sig") as f:   # -sig: PowerShell wrote a BOM
                prev = json.load(f)
        except Exception:  # noqa: BLE001
            prev = None

    # registry (fatal if unreadable)
    reg = os.path.join(src, "registry", REGISTRY_FILE)
    try:
        wb = openpyxl.load_workbook(reg, data_only=True)
        settings_rows = read_tab(wb, "Settings")
        bell = read_tab(wb, "Bell")
        daytypes = read_tab(wb, "DayTypes")
        rooms = read_tab(wb, "Rooms")
    except Exception as e:  # noqa: BLE001
        msg = "FATAL: cannot read registry '%s' - %s. Nothing published; previous data.json untouched." % (reg, e)
        with open(report_path, "w", encoding="utf-8") as f:
            f.write(msg + "\n")
        print(msg)
        sys.exit(1)
    # src/day-types.txt: date overrides typed from a phone ("2026-10-30  Half Day  Pep assembly").
    # They go first so an exact-date line wins over the DayTypes tab and the weekday rule.
    extra_days = []
    dt_file = os.path.join(src, "day-types.txt")
    if os.path.exists(dt_file):
        with open(dt_file, encoding="utf-8-sig") as f:
            for n, line in enumerate(f, 1):
                line = line.strip()
                if not line or line.startswith("#"):
                    continue
                m = re.match(r"^(\S+)\s+(Regular|Wednesday|Half Day|No School|[A-Za-z][A-Za-z ]*?)(?:\s{2,}|\t|\s+-\s+|$)(.*)$", line)
                try:
                    if not m:
                        raise ValueError("expected DATE, two spaces, DAY TYPE")
                    extra_days.append({"date": parse_date(m.group(1)).isoformat(), "day_type": m.group(2).strip(), "note": m.group(3).strip()})
                except ValueError as e:
                    issues.append("day-types.txt line %d: %s (%s)." % (n, e, line))
    if extra_days:
        known = {str(b.get("day_type", "")).strip().lower() for b in bell} | {"no school"}
        for x in extra_days:
            if x["day_type"].lower() not in known:
                issues.append("day-types.txt: '%s' is not a bell in the registry (have: %s) - screens will treat %s as announcements-only." % (x["day_type"], ", ".join(sorted(known)), x["date"]))
        daytypes = extra_days + daytypes
        notes.append("Day-type overrides from day-types.txt: %s." % "; ".join("%s=%s" % (x["date"], x["day_type"]) for x in extra_days))

    settings = {}
    for r in settings_rows:
        if r.get("key"):
            settings[r["key"]] = r.get("value", "")
    week_starts = {}
    try:
        for w in read_tab(wb, "Weeks"):
            if w.get("week") and w.get("starts"):
                week_starts[int(float(w["week"]))] = parse_date(w["starts"]).isoformat()
    except Exception as e:  # noqa: BLE001
        issues.append("Weeks tab unreadable (%s) - week numbers fall back to counting from first_day." % e)
    anchors = {"q1": "2026-08-19", "q2": "2026-10-26", "q3": "2027-01-19", "q4": "2027-03-22"}
    for k in ("q1", "q2", "q3", "q4"):
        if settings.get(k + "_start"):
            try:
                anchors[k] = parse_date(settings[k + "_start"]).isoformat()
            except ValueError:
                issues.append("Settings: %s_start '%s' is unreadable." % (k, settings[k + "_start"]))
    anchors["s1"], anchors["s2"] = anchors["q1"], anchors["q3"]
    first_day = None
    if settings.get("first_day"):
        try:
            first_day = parse_date(settings["first_day"])
        except ValueError:
            issues.append("Settings: first_day '%s' is unreadable." % settings["first_day"])

    sns_index, sns_notes = build_sns_index(os.path.join(src, "sns"))
    issues.extend(sns_notes)
    sns_used = set()

    # course -> optional content source (first non-blank wins)
    needs, src_map = [], {}
    for r in rooms:
        c = r.get("course", "")
        if not c:
            continue
        if c not in src_map:
            needs.append(c)
            src_map[c] = ""
        if r.get("content_source") and not src_map[c]:
            src_map[c] = r["content_source"]

    today = started.date()
    courses, no_source = {}, []
    for course in sorted(needs):
        source = src_map[course]
        rows_out = None
        key = normalize_course(course)
        if not source:
            hit = key if key in sns_index else None
            if not hit:
                for k in sns_index:
                    if k.startswith(key) or key.startswith(k):
                        hit = k
                        break
            if not hit:
                reg_words = course_words(course)
                cands = [k for k in sns_index if course_words(sns_index[k]["label"]) and all(w in reg_words for w in course_words(sns_index[k]["label"]))]
                if len(cands) == 1:
                    hit = cands[0]
                elif len(cands) > 1:
                    issues.append("Course '%s': several sns files could be it (%s) - make one file's course name exact." % (course, "; ".join(os.path.basename(sns_index[c]["file"]) for c in cands)))
            if hit:
                source = sns_index[hit]["file"]
                sns_used.add(hit)
        if source:
            file = source if os.path.isabs(source) else os.path.join(src, source)
            if not os.path.exists(file):
                issues.append("Missing scope & sequence file: %s (course '%s')." % (source, course))
            else:
                try:
                    rows_out = read_sns_course(file, week_starts, first_day, anchors)
                except Exception as e:  # noqa: BLE001
                    issues.append("S&S file '%s' unreadable (%s)." % (os.path.basename(file), e))
                    rows_out = None
        else:
            no_source.append(course)

        if rows_out is not None:
            src_name = os.path.basename(source)
            active, x_units, cal_unit, cal_date = False, [], "", None
            for u in rows_out:
                if not u["unit"]:
                    continue
                if test_current(u["current"]):
                    active = True
                    x_units.append(u["unit"])
                d = None
                if u["starts"]:
                    try:
                        d = parse_date(u["starts"])
                    except ValueError:
                        issues.append("%s: unit '%s' has an unreadable Starts date '%s'." % (src_name, u["unit"], u["starts"]))
                elif not test_current(u["current"]):
                    issues.append("%s: unit '%s' can't be scheduled - add its week to the registry Weeks tab, use a date, or mark it current." % (src_name, u["unit"]))
                if d and d <= today:
                    active = True
                    if not cal_date or d > cal_date:
                        cal_date, cal_unit = d, u["unit"]
                if (test_current(u["current"]) or (d and d <= today)) and not u["learning_target_s"]:
                    issues.append("%s: unit '%s' is live but has no learning targets." % (src_name, u["unit"]))
            if len(x_units) > 1:
                issues.append("%s: %d units are marked current - the wall shows the last one, '%s'. Clear the extras." % (src_name, len(x_units), x_units[-1]))
            if x_units and cal_unit and x_units[-1] != cal_unit:
                notes.append("%s: pacing override active - showing '%s' while the calendar says '%s'." % (src_name, x_units[-1], cal_unit))
            if not active:
                issues.append("%s: no unit is active today (check dates/weeks or mark one current)." % src_name)
            courses[course] = rows_out
        else:
            if prev and isinstance(prev.get("courses"), dict) and prev["courses"].get(course):
                courses[course] = prev["courses"][course]
                if source:
                    notes.append("'%s': kept yesterday's data on the wall." % course)
            else:
                courses[course] = []

    if no_source:
        shown = "; ".join(no_source[:8]) + ("; ... %d more" % (len(no_source) - 8) if len(no_source) > 8 else "")
        notes.append("%d course(s) have no scope & sequence file yet - their walls show the announcements deck. (%s)" % (len(no_source), shown))
    for k, v in sns_index.items():
        if k not in sns_used:
            issues.append("sns/%s: course name '%s' does not match any Course in the registry - fix the name in the file or the Rooms tab." % (os.path.basename(v["file"]), v["label"]))

    # announcements
    announcements = []
    try:
        if args.html:
            with open(args.html, encoding="utf-8", errors="replace") as f:
                html = f.read()
        elif args.no_fetch:
            raise RuntimeError("fetch skipped (--no-fetch)")
        else:
            html = fetch_announcements(ANNOUNCEMENTS_URL)
        announcements = parse_announcements(html)
        if announcements:
            notes.append("Announcements: %d items from the school site (%s)." % (len(announcements), announcements[0]["date"]))
        else:
            raise RuntimeError("page fetched but no items parsed - layout may have changed")
    except Exception as e:  # noqa: BLE001
        issues.append("Announcements scrape failed (%s) - keeping yesterday's." % e)
        if prev and prev.get("announcements"):
            announcements = list(prev["announcements"])

    # slim the PUBLISHED copy: only what a screen renders (no section numbers,
    # enrollment, teacher numbers, userids or file paths - those stay in the registry)
    public_rooms = [{
        "room": r.get("room", ""), "period": r.get("period", ""), "course": r.get("course", ""),
        "course_title": r.get("course_title", ""), "teacher": honorific(r.get("teacher", "")), "day_type": r.get("day_type", ""),
    } for r in rooms]
    bell_pub = [{k: v for k, v in b.items() if k in ("day_type", "period", "label", "start", "end")} for b in bell]
    dt_pub = [{k: v for k, v in d.items() if k in ("date", "day_type", "note")} for d in daytypes]

    body = {"settings": settings, "bell": bell_pub, "daytypes": dt_pub, "rooms": public_rooms, "courses": courses, "announcements": announcements}

    # Only bump the version when something on the walls changed, or once a day
    # regardless (screens flag data older than 36 h as stale).
    prev_body = {k: prev.get(k) for k in body} if prev else None
    changed = prev_body != body
    prev_gen = None
    if prev and prev.get("generated"):
        try:
            prev_gen = dt.datetime.strptime(prev["generated"][:19], "%Y-%m-%dT%H:%M:%S")
        except ValueError:
            prev_gen = None
    daily_due = (prev_gen is None) or ((started.replace(tzinfo=None) - prev_gen) > dt.timedelta(hours=20))
    if changed or daily_due:
        out = {"generated": started.strftime("%Y-%m-%dT%H:%M:%S")}
        out.update(body)
        with open(data_path, "w", encoding="utf-8") as f:
            json.dump(out, f, ensure_ascii=False, indent=1)
        with open(os.path.join(site, "version.txt"), "w", encoding="ascii") as f:
            f.write(out["generated"])
        notes.append("data.json written (%s)." % ("content changed" if changed else "daily refresh, no content change"))
        rc_changed = True
    else:
        notes.append("No changes since last compile - data.json left as is.")
        rc_changed = False

    # report (committed beside the site so it can be read on the phone)
    report = ["EHS Displays compile - %s" % started.strftime("%a %b %d, %Y %-I:%M %p %Z"),
              "Rooms: %d rows | Courses compiled: %d | Issues: %d" % (len(rooms), len(courses), len(issues)), ""]
    if issues:
        report += ["ISSUES (fix in the named file; the wall shows last good data meanwhile):"] + issues
    else:
        report += ["No issues. Every course compiled clean."]
    report += [""] + notes
    with open(report_path, "w", encoding="utf-8") as f:
        f.write("\n".join(report) + "\n")
    # teachers never read compile-report.txt, so the sns lines also land where they look
    sns_dir = os.path.join(src, "sns")
    sns_note = os.path.join(sns_dir, "_READ ME - problems found last night.txt")
    sns_issues = [i for i in issues if i.startswith("sns/")]
    if sns_issues and os.path.isdir(sns_dir):
        with open(sns_note, "w", encoding="utf-8") as f:
            f.write("Checked %s. Each line names the file and what to fix; your wall shows the last good version until then.\n\n" % started.strftime("%a %b %d, %-I:%M %p"))
            f.write("\n".join(sns_issues) + "\n")
    elif os.path.exists(sns_note):
        os.remove(sns_note)
    print("\n".join(report))
    # exit code 0 either way; the workflow decides whether there is anything to commit
    with open(os.path.join(site, ".compile-changed"), "w") as f:
        f.write("1" if rc_changed else "0")


if __name__ == "__main__":
    main()
